"""
Load company masterdata (Excel/CSV) into public.roc_companies for instant CIN checks.
"""

from __future__ import annotations

import csv
import logging
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

logger = logging.getLogger("fundx.company_masterdata")

EXPECTED_HEADERS = {
    "cin",
    "company_name",
    "registration_date",
    "category",
    "company_class",
    "listing_status",
    "authorized_capital",
    "paidup_capital",
    "roc",
    "address",
    "state",
    "company_status",
    "industrial_classification",
}


def _masterdata_paths() -> list[Path]:
    env = os.getenv("COMPANY_MASTERDATA_PATH", "").strip()
    candidates: list[Path] = []
    if env:
        candidates.append(Path(env))
    # Docker: ./database mounted at /app/database
    candidates.extend(
        [
            Path("/app/database/masterdata/company_masterdata.xlsx"),
            Path("/app/database/masterdata/company_masterdata.csv"),
            Path(__file__).resolve().parents[2] / "database" / "masterdata" / "company_masterdata.xlsx",
            Path(__file__).resolve().parents[2] / "database" / "masterdata" / "company_masterdata.csv",
            Path("database/masterdata/company_masterdata.xlsx"),
            Path("database/masterdata/company_masterdata.csv"),
        ]
    )
    seen: set[str] = set()
    out: list[Path] = []
    for p in candidates:
        key = str(p.resolve()) if p.exists() else str(p)
        if key in seen:
            continue
        seen.add(key)
        out.append(p)
    return out


def _parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _normalize_row(row: dict[str, Any]) -> dict[str, Any] | None:
    cin = str(row.get("cin") or "").strip().upper()
    name = str(row.get("company_name") or "").strip()
    if not cin or not name:
        return None
    status = str(row.get("company_status") or "Active").strip() or "Active"
    return {
        "cin": cin,
        "company_name": name,
        "registration_date": _parse_date(row.get("registration_date")),
        "category": (str(row.get("category") or "").strip() or None),
        "company_class": (str(row.get("company_class") or "").strip() or None),
        "listing_status": (str(row.get("listing_status") or "").strip() or None),
        "authorized_capital": _parse_number(row.get("authorized_capital")),
        "paidup_capital": _parse_number(row.get("paidup_capital")),
        "roc": (str(row.get("roc") or "").strip() or None),
        "address": (str(row.get("address") or "").strip() or None),
        "state": (str(row.get("state") or "").strip() or None),
        "company_status": status,
        "industrial_classification": (
            str(row.get("industrial_classification") or "").strip() or None
        ),
    }


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        return [r for r in reader if r]


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to read company_masterdata.xlsx — pip install openpyxl"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers_raw = next(rows_iter, None)
    if not headers_raw:
        return []
    headers = [str(h or "").strip().lower() for h in headers_raw]
    out: list[dict[str, Any]] = []
    for values in rows_iter:
        row = {headers[i]: values[i] for i in range(len(headers)) if i < len(values)}
        out.append(row)
    return out


def load_masterdata_rows() -> tuple[Path | None, list[dict[str, Any]]]:
    for path in _masterdata_paths():
        if not path.exists():
            continue
        try:
            raw = _read_xlsx(path) if path.suffix.lower() in {".xlsx", ".xlsm"} else _read_csv(path)
        except Exception as exc:  # noqa: BLE001
            logger.warning("Failed reading masterdata %s: %s", path, exc)
            continue
        normalized = []
        for row in raw:
            item = _normalize_row({str(k).strip().lower(): v for k, v in row.items()})
            if item:
                normalized.append(item)
        if normalized:
            return path, normalized
    return None, []


UPSERT_SQL = """
INSERT INTO public.roc_companies (
  cin, company_name, registration_date, category, company_class, listing_status,
  authorized_capital, paidup_capital, roc, address, state, company_status,
  industrial_classification
) VALUES (
  $1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13
)
ON CONFLICT (cin) DO UPDATE SET
  company_name = EXCLUDED.company_name,
  registration_date = EXCLUDED.registration_date,
  category = EXCLUDED.category,
  company_class = EXCLUDED.company_class,
  listing_status = EXCLUDED.listing_status,
  authorized_capital = EXCLUDED.authorized_capital,
  paidup_capital = EXCLUDED.paidup_capital,
  roc = EXCLUDED.roc,
  address = EXCLUDED.address,
  state = EXCLUDED.state,
  company_status = EXCLUDED.company_status,
  industrial_classification = EXCLUDED.industrial_classification,
  updated_at = NOW()
"""


async def seed_roc_companies_from_masterdata(pool) -> dict[str, Any]:
    """Upsert Excel/CSV company masterdata into roc_companies."""
    if pool is None:
        return {"seeded": 0, "path": None, "error": "no_db_pool"}

    path, rows = load_masterdata_rows()
    if not rows:
        logger.warning("No company masterdata file found — inserting SQL fallback demo CINs")
        rows = [
            {
                "cin": "U40100WB2022PTC256639",
                "company_name": "AeroGrid Tech Innovations Private Limited",
                "registration_date": date(2022, 10, 15),
                "category": "Private",
                "company_class": "Private Company",
                "listing_status": "Unlisted",
                "authorized_capital": 5000000,
                "paidup_capital": 2500000,
                "roc": "ROC-KOLKATA",
                "address": "42/A Park Lane, Ballygunge, Kolkata 700019",
                "state": "West Bengal",
                "company_status": "Active",
                "industrial_classification": "Renewable Energy / Climate Technology",
            },
            {
                "cin": "U72900KA2024PTC182001",
                "company_name": "NovaGrid Energy Systems Private Limited",
                "registration_date": date(2024, 3, 14),
                "category": "Private",
                "company_class": "Private Company",
                "listing_status": "Unlisted",
                "authorized_capital": 1000000,
                "paidup_capital": 100000,
                "roc": "ROC-BANGALORE",
                "address": "Electronic City Phase 1, Bengaluru 560100",
                "state": "Karnataka",
                "company_status": "Active",
                "industrial_classification": "CleanTech / Energy Storage",
            },
            {
                "cin": "U35201WB2023PTC262723",
                "company_name": "Legacy Commerce Enterprise Limited",
                "registration_date": date(2023, 2, 14),
                "category": "Private",
                "company_class": "Private Company",
                "listing_status": "Unlisted",
                "authorized_capital": 500000,
                "paidup_capital": 250000,
                "roc": "ROC-KOLKATA",
                "address": "89 Lower Circular Road, Kolkata 700013",
                "state": "West Bengal",
                "company_status": "Strike Off",
                "industrial_classification": "General Trading",
            },
        ]
        path = None

    count = 0
    async with pool.acquire() as conn:
        for row in rows:
            await conn.execute(
                UPSERT_SQL,
                row["cin"],
                row["company_name"],
                row["registration_date"],
                row["category"],
                row["company_class"],
                row["listing_status"],
                row["authorized_capital"],
                row["paidup_capital"],
                row["roc"],
                row["address"],
                row["state"],
                row["company_status"],
                row["industrial_classification"],
            )
            count += 1

    logger.info("Seeded %s ROC companies from %s", count, path or "sql-fallback")
    return {"seeded": count, "path": str(path) if path else "sql-fallback"}
